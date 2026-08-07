"""docs/ の項目一覧 xlsx を全シート TSV としてダンプする一時スクリプト。"""
import sys
from pathlib import Path
import openpyxl

src = Path(sys.argv[1])
outdir = Path(sys.argv[2])
outdir.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
print("SHEETS:", wb.sheetnames)

index = []
for name in wb.sheetnames:
    ws = wb[name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).replace("\t", " ").replace("\n", " / ") for c in row]
        while cells and cells[-1] == "":
            cells.pop()
        rows.append("\t".join(cells))
    while rows and rows[-1].strip() == "":
        rows.pop()
    safe = "".join(ch if ch.isalnum() or ch in "-_一二三四五六七八九十" else "_" for ch in name)
    path = outdir / f"{safe}.tsv"
    path.write_text("\n".join(rows), encoding="utf-8")
    index.append((name, len(rows), str(path)))

for name, n, path in index:
    print(f"{name}\t{n} rows\t{path}")
